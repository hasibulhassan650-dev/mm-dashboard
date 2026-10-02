"""The exported workbooks, built and then read back.

A spreadsheet leaves the building. Once it is in someone's inbox nothing on the
site can explain it, so the file has to carry its own caveats and must not
contain a number that is really an absence. The rule these tests enforce:

    an unknown value is an EMPTY CELL, never 0

because a zero gets summed. "BB has not published an auction calendar for this
month" totalled as zero becomes "no borrowing that month", which is the exact
inversion of the truth.
"""
import datetime
import io
import os
import sys

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.append(os.path.join(_ROOT, "api", "routers"))

from openpyxl import load_workbook

from export_logic import MONEY, build_workbook, filename  # noqa: E402


def _read(body: bytes):
    return load_workbook(io.BytesIO(body))


def _rows(ws):
    return list(ws.iter_rows(values_only=True))


class TestFormatting:
    @pytest.fixture
    def wb(self):
        body = build_workbook(
            "Test report",
            [("Monthly by product", [
                {"Month": "2026-10", "Redemption T-Bond": 5126.24, "Auction Total": 6030.0,
                 "Coupon Rate %": 10.4085, "Months": 3},
                {"Month": "2026-11", "Redemption T-Bond": 11367.0, "Auction Total": None,
                 "Coupon Rate %": None, "Months": 2},
            ])],
            facts=[("Unit", "BDT crore"), ("Window", "2026-10-01 to 2026-11-30")],
            caveats=["Blank means unknown, not zero."],
        )
        return _read(body)

    def test_the_cover_sheet_comes_first_and_carries_provenance(self, wb):
        assert wb.sheetnames[0] == "About this export"
        flat = " ".join(str(c) for row in _rows(wb["About this export"]) for c in row if c)
        assert "BDT crore" in flat
        assert "Generated (UTC)" in flat
        assert "Blank means unknown" in flat

    def test_the_header_row_is_bold_and_frozen(self, wb):
        ws = wb["Monthly by product"]
        assert ws.cell(row=1, column=1).font.bold is True
        # freeze below the header and right of the first column
        assert ws.freeze_panes == "B2"

    def test_money_columns_carry_a_thousands_format(self, wb):
        ws = wb["Monthly by product"]
        assert ws.cell(row=2, column=2).number_format == MONEY

    def test_a_rate_column_is_not_formatted_as_money(self, wb):
        ws = wb["Monthly by product"]
        hdrs = [c.value for c in ws[1]]
        col = hdrs.index("Coupon Rate %") + 1
        assert ws.cell(row=2, column=col).number_format == "0.0000"

    def test_columns_are_given_a_usable_width(self, wb):
        ws = wb["Monthly by product"]
        assert all(d.width >= 11 for d in ws.column_dimensions.values())

    def test_an_autofilter_is_set(self, wb):
        assert wb["Monthly by product"].auto_filter.ref is not None

    def test_an_empty_sheet_says_so_rather_than_being_blank(self):
        wb = _read(build_workbook("T", [("Nothing", [])]))
        assert "No rows" in str(_rows(wb["Nothing"])[0][0])


class TestUnknownIsBlankNotZero:
    def test_a_none_stays_an_empty_cell(self):
        wb = _read(build_workbook("T", [("S", [
            {"Month": "2027-01", "Auction Total": None},
            {"Month": "2026-10", "Auction Total": 6030.0},
        ])]))
        ws = wb["S"]
        hdrs = [c.value for c in ws[1]]
        col = hdrs.index("Auction Total") + 1
        assert ws.cell(row=2, column=col).value is None, "unknown must not become 0"
        assert ws.cell(row=3, column=col).value == 6030.0

    def test_a_real_zero_is_still_written(self):
        # The rule is about ABSENCE. A genuine zero -- a T-Bill coupon, which is
        # zero by construction -- must survive as 0.
        wb = _read(build_workbook("T", [("S", [{"Coupon T-Bill": 0.0}])]))
        assert wb["S"].cell(row=2, column=1).value == 0

    def test_a_ragged_row_set_keeps_every_column(self):
        # Later rows carrying keys the first row lacks must not be dropped --
        # that would silently lose a product from the export.
        wb = _read(build_workbook("T", [("S", [
            {"A": 1.0}, {"A": 2.0, "B": 3.0},
        ])]))
        assert [c.value for c in wb["S"][1]] == ["A", "B"]
        assert wb["S"].cell(row=2, column=2).value is None


class TestFilename:
    def test_parts_are_joined_and_spaces_removed(self):
        assert filename("bd_gsec_schedule", "20y", "2026-10-01") == \
            "bd_gsec_schedule_20y_2026-10-01.xlsx"

    def test_empty_parts_are_skipped(self):
        assert filename("x", None, "", "a") == "x_a.xlsx"


class TestSheetNames:
    def test_an_illegal_sheet_name_is_sanitised(self):
        # Excel refuses []:*?/\ and anything over 31 chars; a rejected name
        # means the whole download fails to open.
        wb = _read(build_workbook("T", [("a/b:c*d?e[f]g" + "x" * 40, [{"A": 1}])]))
        name = wb.sheetnames[1]
        assert len(name) <= 31
        assert not any(c in name for c in "[]:*?/\\")


class TestTheDeploymentBoundary:
    def test_export_logic_imports_nothing_from_the_repo_root(self):
        import ast

        import export_logic
        forbidden = {"config", "db", "engines", "calendar_utils", "validate",
                     "fetchers", "parsers"}
        src = io.open(export_logic.__file__, encoding="utf-8").read()
        bad = []
        for node in ast.walk(ast.parse(src)):
            if isinstance(node, ast.Import):
                bad += [a.name for a in node.names if a.name.split(".")[0] in forbidden]
            elif isinstance(node, ast.ImportFrom) and node.level == 0 and node.module:
                if node.module.split(".")[0] in forbidden:
                    bad.append(node.module)
        assert not bad, f"export_logic cannot import from the repo root: {bad}"

    def test_openpyxl_is_in_the_apis_own_requirements(self):
        # The repo root has it, but the API deploys rooted at api/ and cannot
        # see that file. Omitting it here is a production-only 500.
        with open(os.path.join(_ROOT, "api", "requirements.txt"), encoding="utf-8") as fh:
            assert "openpyxl" in fh.read()
